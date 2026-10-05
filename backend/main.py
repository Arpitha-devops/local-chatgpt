import io
import json
import os
from typing import List, Optional

import httpx
from fastapi import FastAPI, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

import db

OLLAMA_URL = os.getenv("OLLAMA_URL", "http://localhost:11434")
MAX_UPLOAD_BYTES = 20 * 1024 * 1024
# Ollama's default context window is small and silently drops what doesn't fit, so chats with
# documents ask for a bigger one, up to this many tokens.
MAX_NUM_CTX = 16384
# Added to chats that have documents so every answer, including follow-ups, comes from them.
DOCUMENT_PROMPT = (
    "The user has uploaded one or more documents, shown inside <document> tags. "
    "Answer every question using only the information in those documents. "
    "If the answer is not in the documents, say that the documents do not contain it; "
    "do not answer from your general knowledge."
)

app = FastAPI(title="Local ChatGPT")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_methods=["*"],
    allow_headers=["*"],
)
db.init()


class Document(BaseModel):
    name: str
    text: str


class Message(BaseModel):
    role: str
    content: str
    documents: Optional[List[Document]] = None


class ChatRequest(BaseModel):
    model: str
    messages: List[Message]
    system: Optional[str] = None


class StoredMessage(Message):
    stats: Optional[dict] = None


class Conversation(BaseModel):
    title: str
    updatedAt: int
    messages: List[StoredMessage]


def usage_stats(data):
    """Token counts and timings from Ollama's final chunk (durations are in nanoseconds)."""
    prompt_tokens = data.get("prompt_eval_count", 0)
    output_tokens = data.get("eval_count", 0)
    eval_s = data.get("eval_duration", 0) / 1e9
    return {
        "model": data.get("model"),
        "prompt_tokens": prompt_tokens,
        "output_tokens": output_tokens,
        "total_tokens": prompt_tokens + output_tokens,
        "total_seconds": round(data.get("total_duration", 0) / 1e9, 2),
        "load_seconds": round(data.get("load_duration", 0) / 1e9, 2),
        "tokens_per_second": round(output_tokens / eval_s, 1) if eval_s else None,
        "cost_usd": 0.0,  # local inference, no per-token charges
    }


@app.get("/api/health")
async def health():
    try:
        async with httpx.AsyncClient(timeout=3) as client:
            r = await client.get(f"{OLLAMA_URL}/api/version")
            r.raise_for_status()
            return {"ok": True, "ollama": r.json().get("version")}
    except httpx.HTTPError:
        return {"ok": False, "ollama": None}


@app.get("/api/models")
async def models():
    try:
        async with httpx.AsyncClient(timeout=10) as client:
            r = await client.get(f"{OLLAMA_URL}/api/tags")
            r.raise_for_status()
    except httpx.HTTPError as e:
        raise HTTPException(502, f"Cannot reach Ollama at {OLLAMA_URL}: {e}")
    return {"models": [m["name"] for m in r.json().get("models", [])]}


@app.get("/api/conversations")
def list_conversations():
    return {"conversations": db.list_conversations()}


@app.put("/api/conversations/{id}")
def save_conversation(id: str, conv: Conversation):
    db.save_conversation(id, conv.title, conv.updatedAt, [m.dict() for m in conv.messages])
    return {"ok": True}


@app.delete("/api/conversations/{id}")
def delete_conversation(id: str):
    db.delete_conversation(id)
    return {"ok": True}


def extract_text(name, data):
    """Plain text of an uploaded PDF, Word or text file."""
    ext = os.path.splitext(name)[1].lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            return "\n\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
        if ext == ".docx":
            import docx

            return "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
    except Exception as e:
        raise HTTPException(422, f"Could not read {name}: {e}")
    if b"\x00" in data:
        raise HTTPException(415, f"{name}: only PDF, Word (.docx) and text files are supported")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise HTTPException(415, f"{name}: only PDF, Word (.docx) and text files are supported")


@app.post("/api/documents")
def upload_document(file: UploadFile):
    name = file.filename or "document"
    data = file.file.read(MAX_UPLOAD_BYTES + 1)
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, f"{name} is larger than {MAX_UPLOAD_BYTES // (1024 * 1024)} MB")
    text = extract_text(name, data).strip()
    if not text:
        raise HTTPException(422, f"No text found in {name} (scanned PDFs are not supported)")
    return {"name": name, "text": text}


def ollama_message(m: Message):
    """Folds a message's attached documents into the text the model sees."""
    content = m.content
    if m.documents:
        docs = "\n\n".join(f'<document name="{d.name}">\n{d.text}\n</document>' for d in m.documents)
        content = f"{docs}\n\n{content}".strip()
    return {"role": m.role, "content": content}


@app.post("/api/chat")
async def chat(req: ChatRequest):
    messages = [ollama_message(m) for m in req.messages]
    has_documents = any(m.documents for m in req.messages)
    system = "\n\n".join(filter(None, [req.system, DOCUMENT_PROMPT if has_documents else None]))
    if system:
        messages.insert(0, {"role": "system", "content": system})
    payload = {"model": req.model, "messages": messages, "stream": True}
    if has_documents:
        # Rough estimate: ~3 characters per token, plus room for the reply.
        needed = sum(len(m["content"]) for m in messages) // 3 + 1024
        if needed > 4096:
            payload["options"] = {"num_ctx": min(needed, MAX_NUM_CTX)}

    # Streams NDJSON events, one per line:
    #   {"type": "token", "content": "..."}
    #   {"type": "done", "stats": {...}}
    #   {"type": "error", "message": "..."}
    def event(obj):
        return json.dumps(obj) + "\n"

    async def stream():
        try:
            async with httpx.AsyncClient(timeout=None) as client:
                async with client.stream("POST", f"{OLLAMA_URL}/api/chat", json=payload) as r:
                    if r.status_code != 200:
                        body = (await r.aread()).decode(errors="ignore")
                        yield event({"type": "error", "message": f"Error from Ollama: {body}"})
                        return
                    async for line in r.aiter_lines():
                        if not line:
                            continue
                        data = json.loads(line)
                        if data.get("error"):
                            yield event({"type": "error", "message": data["error"]})
                            return
                        chunk = data.get("message", {}).get("content", "")
                        if chunk:
                            yield event({"type": "token", "content": chunk})
                        if data.get("done"):
                            yield event({"type": "done", "stats": usage_stats(data)})
                            return
        except httpx.HTTPError as e:
            yield event({"type": "error", "message": f"Cannot reach Ollama at {OLLAMA_URL}: {e}"})

    return StreamingResponse(
        stream(),
        media_type="application/x-ndjson",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )
